"""存储层：SQLite 连接与 schema、Excel 读写、users.json、版本快照、备份、settings。

与原版保持完全一致的 SQLite schema 和 Excel 表头/目录结构，
生产切换时可直接复用原版 data 目录（纯文件复制，无迁移脚本）。

差异：砍掉了 settings 回写 app.ini 的双向同步，settings 只以数据库为准。
"""
from __future__ import annotations

import csv
import io
import json
import logging
import shutil
import sqlite3
import subprocess
import threading
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from . import notifier
from .core import (
    ENTITIES,
    clean_value,
    entity_meta,
    normalize_entity,
    normalize_row,
    now_text,
    pending_status_text,
)

logger = logging.getLogger(__name__)

LOCK = threading.RLock()
NOTIFY_TIMEOUT = 5  # 后台通知发送超时（秒），覆盖 TCP 连接/读取阶段

# 运行期状态：由 create_app() -> init_runtime() 装载，单进程部署下全局唯一。
STATE = {
    "config": None,
    "data_dir": None,
    "official_dir": None,
    "temp_dir": None,
    "versions_dir": None,
    "backups_dir": None,
    "db_path": None,
    "pending_path": None,
}


def init_runtime(config):
    STATE["config"] = config
    data_dir = Path(config["data_dir"])
    STATE["data_dir"] = data_dir
    STATE["official_dir"] = data_dir / "official"
    STATE["temp_dir"] = data_dir / "temp"
    STATE["versions_dir"] = data_dir / "versions"
    STATE["backups_dir"] = data_dir / "backups"
    STATE["db_path"] = data_dir / "app.db"
    STATE["pending_path"] = STATE["temp_dir"] / "待审核变更.xlsx"


def get_users():
    return STATE["config"]["users"]


def get_user_by_token(token):
    token = (token or "").strip()
    if not token:
        return None
    for user in STATE["config"]["users"].values():
        if user.get("api_token") and user["api_token"] == token:
            return user
    return None


class ClosingConnection(sqlite3.Connection):
    def __exit__(self, exc_type, exc_value, traceback):
        result = super().__exit__(exc_type, exc_value, traceback)
        self.close()
        return result


def db():
    conn = sqlite3.connect(STATE["db_path"], factory=ClosingConnection)
    conn.row_factory = sqlite3.Row
    return conn


CHANGES_SCHEMA = """CREATE TABLE IF NOT EXISTS changes (
    id TEXT PRIMARY KEY, entity TEXT NOT NULL, operation TEXT NOT NULL,
    record_key TEXT NOT NULL, new_data TEXT,
    status TEXT NOT NULL, submitter TEXT NOT NULL, submitted_at TEXT NOT NULL,
    reviewer TEXT, reviewed_at TEXT, reject_reason TEXT
)"""

RACK_REQUESTS_SCHEMA = """CREATE TABLE IF NOT EXISTS rack_requests (
    id TEXT PRIMARY KEY,
    serial_number TEXT NOT NULL,
    request_type TEXT NOT NULL,
    request_data TEXT NOT NULL,
    status TEXT NOT NULL,
    submitter_id TEXT NOT NULL,
    submitted_at TEXT NOT NULL,
    stage TEXT NOT NULL DEFAULT '',
    stage_data TEXT NOT NULL DEFAULT '{}',
    allocation_reviewer_id TEXT,
    allocation_at TEXT,
    implementation_reviewer_id TEXT,
    implementation_at TEXT,
    reviewer_id TEXT,
    reviewed_at TEXT,
    reject_reason TEXT
)"""

SETTINGS_SCHEMA = """CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
)"""


def ensure_changes_schema(conn):
    conn.execute(CHANGES_SCHEMA)
    columns = [row["name"] for row in conn.execute("PRAGMA table_info(changes)").fetchall()]
    if "old_data" not in columns:
        return
    conn.execute("ALTER TABLE changes RENAME TO changes_with_old_data")
    conn.execute(CHANGES_SCHEMA)
    conn.execute("""INSERT INTO changes
        (id,entity,operation,record_key,new_data,status,submitter,submitted_at,reviewer,reviewed_at,reject_reason)
        SELECT id,entity,operation,record_key,new_data,status,submitter,submitted_at,reviewer,reviewed_at,reject_reason
        FROM changes_with_old_data""")
    conn.execute("DROP TABLE changes_with_old_data")


def add_column_if_missing(conn, table_name, column_name, column_sql):
    columns = [row["name"] for row in conn.execute(f"PRAGMA table_info({table_name})").fetchall()]
    if column_name not in columns:
        conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_sql}")


def ensure_rack_requests_schema(conn):
    from .core import initial_rack_request_stage

    conn.execute(RACK_REQUESTS_SCHEMA)
    add_column_if_missing(conn, "rack_requests", "stage", "stage TEXT NOT NULL DEFAULT ''")
    add_column_if_missing(conn, "rack_requests", "stage_data", "stage_data TEXT NOT NULL DEFAULT '{}'")
    add_column_if_missing(conn, "rack_requests", "allocation_reviewer_id", "allocation_reviewer_id TEXT")
    add_column_if_missing(conn, "rack_requests", "allocation_at", "allocation_at TEXT")
    add_column_if_missing(conn, "rack_requests", "implementation_reviewer_id", "implementation_reviewer_id TEXT")
    add_column_if_missing(conn, "rack_requests", "implementation_at", "implementation_at TEXT")
    conn.execute("""UPDATE rack_requests
        SET stage = CASE
            WHEN request_type='上架' THEN 'allocation'
            WHEN request_type='下架' THEN 'implementation'
            ELSE stage
        END
        WHERE status='pending' AND (stage IS NULL OR stage='')""")
    conn.execute("UPDATE rack_requests SET stage='completed' WHERE status='approved' AND (stage IS NULL OR stage='')")
    conn.execute("UPDATE rack_requests SET stage='rejected' WHERE status='rejected' AND (stage IS NULL OR stage='')")
    conn.execute("UPDATE rack_requests SET stage_data='{}' WHERE stage_data IS NULL OR stage_data=''")


def ensure_default_settings(conn):
    config = STATE["config"]
    config_defaults = {
        "backup_keep_count": str(config["backup_keep_count"]),
        "backup_command": config["backup_command"],
        "audit_keep_count": str(config["audit_keep_count"]),
        "notify_enabled": str(config["notify_enabled"]),
        "notify_webhook_url": config["notify_webhook_url"],
        "notify_message_type": config["notify_message_type"],
    }
    reset = config["reset"]
    for key, value in config_defaults.items():
        if reset:
            conn.execute("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)", (key, value))
        else:
            conn.execute("INSERT OR IGNORE INTO settings (key,value) VALUES (?,?)", (key, value))
    if reset:
        config["reset"] = 0
    conn.execute("INSERT OR IGNORE INTO settings (key,value) VALUES (?,?)", ("last_backup_at", ""))


def read_settings(conn=None):
    owns = conn is None
    conn = conn or db()
    try:
        ensure_default_settings(conn)
        data = {row["key"]: row["value"] for row in conn.execute("SELECT key,value FROM settings")}
        return {
            "backup_keep_count": int(data.get("backup_keep_count") or 0),
            "backup_command": data.get("backup_command", ""),
            "audit_keep_count": int(data.get("audit_keep_count") or 0),
            "notify_enabled": int(data.get("notify_enabled") or 0),
            "notify_webhook_url": data.get("notify_webhook_url", ""),
            "notify_message_type": data.get("notify_message_type", "markdown") or "markdown",
            "last_backup_at": data.get("last_backup_at", ""),
        }
    finally:
        if owns:
            conn.close()


def write_settings(values, conn=None):
    allowed = {
        "backup_keep_count",
        "backup_command",
        "audit_keep_count",
        "notify_enabled",
        "notify_webhook_url",
        "notify_message_type",
        "last_backup_at",
    }
    owns = conn is None
    conn = conn or db()
    try:
        ensure_default_settings(conn)
        for key in allowed:
            if key in values:
                conn.execute("INSERT OR REPLACE INTO settings (key,value) VALUES (?,?)", (key, clean_value(values[key])))
        if owns:
            conn.commit()
        return read_settings(conn)
    finally:
        if owns:
            conn.close()


def parse_notify_enabled(value):
    return 1 if clean_value(value).lower() in {"1", "true", "yes", "on", "启用"} else 0


def validate_notification_settings(values):
    message_type = clean_value(values.get("notify_message_type") or "markdown").lower()
    if message_type not in {"text", "markdown"}:
        raise ValueError("通知格式只能是 text 或 markdown")
    webhook_url = clean_value(values.get("notify_webhook_url"))
    if webhook_url and not webhook_url.startswith(("http://", "https://")):
        raise ValueError("通知 Webhook URL 必须以 http:// 或 https:// 开头")
    return {
        "notify_enabled": parse_notify_enabled(values.get("notify_enabled")),
        "notify_webhook_url": webhook_url,
        "notify_message_type": message_type,
    }


def create_empty_workbook(meta):
    path = STATE["official_dir"] / meta["filename"]
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = meta["sheet"]
    ws.append(meta["headers"])
    wb.save(path)
    wb.close()
    logger.info("created official workbook path=%s sheet=%s", path, meta["sheet"])


def copy_cell_style(source, target):
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format


def read_records(entity):
    entity = normalize_entity(entity)
    meta = ENTITIES[entity]
    path = STATE["official_dir"] / meta["filename"]
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[meta["sheet"]] if meta["sheet"] in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [clean_value(v) for v in rows[0]]
        result = []
        for values in rows[1:]:
            item = {h: clean_value(v) for h, v in zip(headers, values) if h}
            normalized = normalize_row(entity, item)
            if any(normalized.values()) and not (entity == "servers" and normalized.get("序列号", "").startswith("示例-")):
                result.append(normalized)
        return result
    finally:
        wb.close()


def write_records(entity, records):
    entity = normalize_entity(entity)
    meta = ENTITIES[entity]
    path = STATE["official_dir"] / meta["filename"]
    if not path.exists():
        create_empty_workbook(meta)
    wb = load_workbook(path)
    try:
        ws = wb[meta["sheet"]] if meta["sheet"] in wb.sheetnames else wb[wb.sheetnames[0]]
        for c_idx, header in enumerate(meta["headers"], 1):
            ws.cell(1, c_idx, header)
        if ws.max_column > len(meta["headers"]):
            ws.delete_cols(len(meta["headers"]) + 1, ws.max_column - len(meta["headers"]))
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        template_row = 2 if ws.max_row >= 2 else None
        for r_idx, item in enumerate(records, 2):
            for c_idx, header in enumerate(meta["headers"], 1):
                cell = ws.cell(r_idx, c_idx, normalize_row(entity, item).get(header, ""))
                if template_row:
                    copy_cell_style(ws.cell(template_row, c_idx), cell)
        wb.save(path)
    finally:
        wb.close()


def workbook_headers(path, sheet):
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        return [clean_value(value) for value in next(ws.iter_rows(max_row=1, values_only=True), ())]
    finally:
        wb.close()


def sync_official_workbook_headers():
    for entity, meta in ENTITIES.items():
        path = STATE["official_dir"] / meta["filename"]
        if workbook_headers(path, meta["sheet"]) != meta["headers"]:
            write_records(entity, read_records(entity))


def sync_pending_excel(conn=None):
    owns = conn is None
    conn = conn or db()
    rows = conn.execute("SELECT * FROM changes WHERE status='pending' ORDER BY submitted_at").fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "待审核变更"
    headers = ["变更编号", "数据类型", "操作", "记录主键", "审批后", "提交人", "提交时间", "状态"]
    ws.append(headers)
    for row in rows:
        ws.append([row["id"], entity_meta(row["entity"])["label"], row["operation"], row["record_key"], row["new_data"] or "", row["submitter"], row["submitted_at"], "待审核"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [16, 12, 10, 24, 55, 14, 22, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    wb.save(STATE["pending_path"])
    if owns:
        conn.close()


def init_storage():
    STATE["official_dir"].mkdir(parents=True, exist_ok=True)
    STATE["temp_dir"].mkdir(parents=True, exist_ok=True)
    STATE["versions_dir"].mkdir(parents=True, exist_ok=True)
    STATE["backups_dir"].mkdir(parents=True, exist_ok=True)
    for meta in ENTITIES.values():
        create_empty_workbook(meta)
    sync_official_workbook_headers()
    with db() as conn:
        ensure_changes_schema(conn)
        ensure_rack_requests_schema(conn)
        conn.execute(SETTINGS_SCHEMA)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_changes_status ON changes(status)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_rack_requests_status ON rack_requests(status)")
        conn.execute("""CREATE TABLE IF NOT EXISTS versions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, version_code TEXT UNIQUE NOT NULL,
            change_id TEXT NOT NULL, entity TEXT NOT NULL, operation TEXT NOT NULL,
            record_key TEXT NOT NULL, reviewer TEXT NOT NULL, created_at TEXT NOT NULL,
            snapshot_dir TEXT NOT NULL
        )""")
        ensure_default_settings(conn)
    sync_pending_excel()
    with db() as conn:
        swept = sweep_orphan_snapshots(conn)
    if swept:
        logger.info("startup swept %d orphan version snapshot(s)", swept)


def export_headers(meta, include_pending_status=False):
    return ["在途状态", *meta["headers"]] if include_pending_status else meta["headers"]


def export_row_values(meta, row, include_pending_status=False):
    values = [row.get(header, "") for header in meta["headers"]]
    return [pending_status_text(row), *values] if include_pending_status else values


def records_to_xlsx(meta, records, include_pending_status=False):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = meta["sheet"]
    ws.append(export_headers(meta, include_pending_status))
    for row in records:
        ws.append(export_row_values(meta, row, include_pending_status))
    wb.save(output)
    wb.close()
    output.seek(0)
    return output


def records_to_csv(meta, records, include_pending_status=False):
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow(export_headers(meta, include_pending_status))
    for row in records:
        writer.writerow(export_row_values(meta, row, include_pending_status))
    return io.BytesIO(("\ufeff" + text.getvalue()).encode("utf-8"))


def worksheet_header_has_merged_cells(ws):
    return any(cell_range.min_row <= 1 <= cell_range.max_row for cell_range in ws.merged_cells.ranges)


def worksheet_rows_with_merged_values(ws):
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    merged_values = {}
    for cell_range in ws.merged_cells.ranges:
        value = ws.cell(cell_range.min_row, cell_range.min_col).value
        for row_index in range(max(2, cell_range.min_row), cell_range.max_row + 1):
            for column_index in range(cell_range.min_col, cell_range.max_col + 1):
                merged_values[(row_index, column_index)] = value
    for row_index, row in enumerate(rows, start=1):
        for column_index in range(1, len(row) + 1):
            merged_value = merged_values.get((row_index, column_index))
            if merged_value is not None:
                row[column_index - 1] = merged_value
    return rows


def uploaded_table_rows(uploaded, meta):
    filename = clean_value(uploaded.filename).lower()
    data = uploaded.read()
    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), data_only=True)
        try:
            ws = wb[meta["sheet"]] if meta["sheet"] in wb.sheetnames else wb[wb.sheetnames[0]]
            if worksheet_header_has_merged_cells(ws):
                raise ValueError("第一行字段表头不能合并单元格")
            return "Excel", worksheet_rows_with_merged_values(ws)
        finally:
            wb.close()
    if filename.endswith(".csv"):
        text = data.decode("utf-8-sig")
        return "CSV", list(csv.reader(io.StringIO(text)))
    return None, None


def users_json_path():
    return STATE["data_dir"] / "users.json"


def system_users_list():
    return [dict(user) for user in STATE["config"]["users"].values()]


def validate_system_users(users):
    from .core import ROLE_ADMIN, ROLE_BUSINESS, ROLE_ROOM

    if not isinstance(users, list) or not users:
        raise ValueError("用户列表不能为空")
    result = []
    seen = set()
    for index, raw in enumerate(users, 1):
        if not isinstance(raw, dict):
            raise ValueError(f"第 {index} 个用户格式无效")
        user = normalize_system_user(raw)
        if not user["用户id"]:
            raise ValueError(f"第 {index} 个用户id不能为空")
        if user["用户id"] in seen:
            raise ValueError(f"用户id {user['用户id']} 重复")
        if not user["密码"]:
            raise ValueError(f"用户 {user['用户id']} 密码不能为空")
        if user["身份"] not in {ROLE_ADMIN, ROLE_ROOM, ROLE_BUSINESS}:
            raise ValueError(f"用户 {user['用户id']} 身份无效")
        if user["身份"] == ROLE_BUSINESS and not user["业务单位"]:
            raise ValueError(f"业务人员 {user['用户id']} 必须填写业务单位")
        if user["身份"] == ROLE_ROOM and not user["所属机房"]:
            raise ValueError(f"机房人员 {user['用户id']} 必须填写所属机房")
        seen.add(user["用户id"])
        result.append(user)
    return result


def normalize_system_user(raw):
    from .config import normalize_user

    return normalize_user(raw)


def save_system_users(users):
    validated = validate_system_users(users)
    path = users_json_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(validated, ensure_ascii=False, indent=2), encoding="utf-8")
    STATE["config"]["users"] = {user["用户id"]: user for user in validated}
    return validated


def backup_sequence():
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    for index in range(1, 10000):
        candidate = STATE["backups_dir"] / f"B{stamp}-{index:04d}"
        if not candidate.exists():
            return candidate
    raise RuntimeError("无法生成备份目录")


def prune_backups(keep_count):
    if keep_count <= 0:
        return
    backups_dir = STATE["backups_dir"]
    backups = sorted([path for path in backups_dir.glob("B*") if path.is_dir()], key=lambda p: p.name, reverse=True)
    for path in backups[keep_count:]:
        resolved_backup = path.resolve()
        resolved_root = backups_dir.resolve()
        try:
            resolved_backup.relative_to(resolved_root)
        except ValueError:
            logger.warning("skip deleting backup outside backups dir path=%s", path)
            continue
        if resolved_backup != resolved_root:
            shutil.rmtree(path, ignore_errors=True)


def run_backup_command(command, cwd):
    if not command:
        return
    subprocess.run(command, cwd=cwd, shell=True, check=True)


def create_backup(reason="manual"):
    settings = read_settings()
    STATE["backups_dir"].mkdir(parents=True, exist_ok=True)
    target = backup_sequence()
    target.mkdir(parents=True, exist_ok=False)
    for entity in ("customers", "cabinets", "servers"):
        meta = ENTITIES[entity]
        source = STATE["official_dir"] / meta["filename"]
        if source.exists():
            shutil.copy2(source, target / meta["filename"])
    for source in (STATE["db_path"], users_json_path()):
        if source.exists():
            shutil.copy2(source, target / source.name)
    (target / "backup.json").write_text(json.dumps({
        "reason": reason,
        "created_at": now_text(),
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    write_settings({"last_backup_at": now_text()})
    try:
        run_backup_command(settings.get("backup_command", ""), STATE["config"]["base_dir"])
    finally:
        prune_backups(int(settings.get("backup_keep_count") or 0))
    return target


def prune_completed_rows(conn, table_name, keep_count):
    if keep_count <= 0:
        return
    rows = conn.execute(f"""
        SELECT id FROM {table_name}
        WHERE status IN ('approved','rejected')
        ORDER BY reviewed_at DESC, submitted_at DESC, id DESC
    """).fetchall()
    for row in rows[keep_count:]:
        conn.execute(f"DELETE FROM {table_name} WHERE id=?", (row["id"],))


def prune_audit_records(conn):
    keep_count = int(read_settings(conn).get("audit_keep_count") or 0)
    prune_completed_rows(conn, "changes", keep_count)
    prune_completed_rows(conn, "rack_requests", keep_count)


def clear_version_snapshots():
    versions_dir = STATE["versions_dir"]
    versions_dir.mkdir(parents=True, exist_ok=True)
    resolved_root = versions_dir.resolve()
    removed = 0
    for path in list(versions_dir.iterdir()):
        try:
            path.resolve().relative_to(resolved_root)
        except ValueError:
            logger.warning("skip deleting version snapshot outside versions dir path=%s", path)
            continue
        if path.resolve() == resolved_root:
            continue
        if path.is_dir():
            shutil.rmtree(path, ignore_errors=True)
        else:
            path.unlink(missing_ok=True)
        removed += 1
    return removed


def reset_flow_state():
    with LOCK, db() as conn:
        counts = {
            "changes": conn.execute("SELECT COUNT(*) FROM changes").fetchone()[0],
            "rack_requests": conn.execute("SELECT COUNT(*) FROM rack_requests").fetchone()[0],
            "versions": conn.execute("SELECT COUNT(*) FROM versions").fetchone()[0],
        }
        conn.execute("DELETE FROM changes")
        conn.execute("DELETE FROM rack_requests")
        conn.execute("DELETE FROM versions")
        removed_snapshots = clear_version_snapshots()
        sync_pending_excel(conn)
    return {**counts, "version_snapshots": removed_snapshots}


def create_version_record(conn, change_id, entity, operation, record_key, reviewer):
    created_at = now_text()
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    sequence = conn.execute("SELECT COALESCE(MAX(id),0)+1 FROM versions").fetchone()[0]
    version_code = f"V{stamp}-{sequence:04d}"
    snapshot_dir = STATE["versions_dir"] / version_code
    snapshot_dir.mkdir(parents=True, exist_ok=False)
    for meta in ENTITIES.values():
        shutil.copy2(STATE["official_dir"] / meta["filename"], snapshot_dir / meta["filename"])
    conn.execute("""INSERT INTO versions
        (version_code,change_id,entity,operation,record_key,reviewer,created_at,snapshot_dir)
        VALUES (?,?,?,?,?,?,?,?)""", (
        version_code, change_id, entity, operation,
        record_key, reviewer, created_at, str(snapshot_dir),
    ))
    prune_versions(conn)
    return version_code


def snapshot_dir_in_versions(snapshot_dir):
    """判断快照目录是否位于当前 versions 目录之内（迁移/换部署路径时可能不在）。"""
    resolved_versions = STATE["versions_dir"].resolve()
    try:
        resolved_snapshot = Path(snapshot_dir).resolve()
    except OSError:
        return False
    return resolved_snapshot != resolved_versions and resolved_versions in resolved_snapshot.parents


def remove_snapshot_dir(snapshot_dir):
    """删除单个快照目录；返回 True 表示目录已不存在（删除成功或本就没有）。

    只允许删除 versions 目录内的目录，避免误删散落在外部的路径。
    删除失败（如权限/占用）不吞异常，返回 False 交由调用方决定是否保留记录。
    """
    path = Path(snapshot_dir)
    if not snapshot_dir_in_versions(path):
        logger.warning("skip deleting version snapshot outside versions dir path=%s", path)
        return False
    try:
        shutil.rmtree(path)
        return True
    except FileNotFoundError:
        return True
    except OSError:
        logger.exception("failed to delete version snapshot path=%s", path)
        return False


def prune_versions(conn):
    max_versions = STATE["config"]["max_versions"]
    if max_versions <= 0:
        return
    rows = conn.execute("SELECT id,snapshot_dir FROM versions ORDER BY id DESC").fetchall()
    for row in rows[max_versions:]:
        snapshot_dir = Path(row["snapshot_dir"])
        if snapshot_dir_in_versions(snapshot_dir):
            # 先删目录：删不掉就保留记录，下次创建版本时重试，避免"库记录已删、目录残留"的孤儿
            if not remove_snapshot_dir(snapshot_dir):
                continue
        else:
            # 记录指向的路径已不在当前 versions 目录（历史迁移遗留），只清记录
            logger.warning("version snapshot not under versions dir, drop record only path=%s", snapshot_dir)
        conn.execute("DELETE FROM versions WHERE id=?", (row["id"],))


def sweep_orphan_snapshots(conn=None):
    """清除 versions 目录下没有任何版本记录引用的快照目录（历史遗留/迁移残留）。

    仅按目录名与 versions 表的 snapshot_dir 匹配，不会误删仍被记录引用的快照。
    """
    owns = conn is None
    conn = conn or db()
    try:
        known = {Path(row["snapshot_dir"]).name for row in conn.execute("SELECT snapshot_dir FROM versions").fetchall()}
        versions_dir = STATE["versions_dir"]
        if not versions_dir.exists():
            return 0
        removed = 0
        for path in sorted(versions_dir.iterdir()):
            if not path.is_dir() or path.name in known:
                continue
            if remove_snapshot_dir(path):
                removed += 1
                logger.info("removed orphan version snapshot path=%s", path)
        return removed
    finally:
        if owns:
            conn.close()
